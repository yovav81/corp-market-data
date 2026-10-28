"""רינדור טבלת ההתראות לתמונה (לטלגרם בנייד) ולקובץ אקסל.

render_png(rows, title, subtitle, out_prefix) -> [נתיבי PNG]   — עד ROWS_PER_IMG שורות לתמונה
render_xlsx(df, path)                                            — טבלה מעוצבת, RTL

הטקסט בעברית עובר bidi (python-bidi) כי PIL מצייר LTR. הגופן: DejaVu Sans (יש בו עברית) מתוך matplotlib.
"""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from PIL import features as _pf

if _pf.check("raqm"):            # PIL עם libraqm מבצע bidi בעצמו — אסור להפעיל get_display (היפוך כפול)
    def get_display(s):  # type: ignore
        return s
else:
    try:
        from bidi.algorithm import get_display
    except Exception:  # noqa: BLE001
        def get_display(s):  # type: ignore
            return s

ROWS_PER_IMG = 25
LINK = {"cpi": "צמוד", "ils": "שקלי", "fx": "מט\"ח"}
# (כותרת, מפתח, רוחב בפיקסלים, יישור) — הסדר מימין לשמאל
COLS = [("סדרה", "name", 270, "r"), ("דירוג", "rating_low", 85, "c"), ("הצמדה", "linkage", 85, "c"), ("מח\"מ", "duration", 75, "c"),
        ("מרווח %", "cur_spread", 100, "c"), ("ממוצע 3ש' %", "mean", 125, "c"), ("+נ\"ב", "dev_bp", 85, "c"), ("Z", "z", 65, "c"),
        ("Z יחסי", "rel_z", 90, "c"), ("אחוזון", "pct_rank", 90, "c"), ("5 ימים נ\"ב", "chg_5d_bp", 110, "c"), ("★ / ⚠ דיווח", "flag", 170, "c")]


def _font(size, bold=False):
    import matplotlib
    d = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    return ImageFont.truetype(str(d / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")), size)


def _fmt(k, v):
    if v is None or (isinstance(v, float) and v != v):
        return "—" if k != "flag" else ""
    if k == "linkage":
        return LINK.get(v, str(v))
    if k in ("cur_spread", "mean"):
        return f"{v:.2f}"
    if k == "duration":
        return f"{v:.1f}"
    if k in ("z", "rel_z"):
        return f"{v:.1f}"
    if k == "pct_rank":
        return f"{v*100:.0f}"
    if k in ("dev_bp", "chg_5d_bp"):
        return f"{v:+.0f}"
    return str(v)


def flag_of(r):
    parts = []
    if isinstance(r.get("covered"), str) and r["covered"]:
        parts.append("★")
    if isinstance(r.get("maya_recent"), str) and r["maya_recent"]:
        parts.append(f"⚠ {r['maya_recent'][8:]}.{r['maya_recent'][5:7]}")
    return " ".join(parts)


def render_png(rows, title, subtitle, out_prefix):
    """rows: רשימת dict/Series עם מפתחות COLS. מחזיר רשימת נתיבים."""
    rows = list(rows)
    paths = []
    f, fb, fh, fs = _font(22), _font(22, bold=True), _font(19, bold=True), _font(16)
    W = sum(c[2] for c in COLS) + 40
    RH, HH, PAD = 40, 48, 20
    chunks = [rows[i:i + ROWS_PER_IMG] for i in range(0, len(rows), ROWS_PER_IMG)] or [[]]
    for ci, chunk in enumerate(chunks):
        sub_lines = [ln.strip() for ln in subtitle.split("|")]
        top = PAD + 40 + 22 * len(sub_lines) + 14
        H = top + HH + RH * max(1, len(chunk)) + PAD
        img = Image.new("RGB", (W, H), "white"); dr = ImageDraw.Draw(img)
        t = title + (f"  ({ci+1}/{len(chunks)})" if len(chunks) > 1 else "")
        dr.text((W - PAD, PAD), get_display(t), font=_font(26, bold=True), fill="#1F3864", anchor="ra")
        for li, ln in enumerate(sub_lines):
            dr.text((W - PAD, PAD + 40 + 22 * li), get_display(ln), font=fs, fill="#555555", anchor="ra")
        y = top
        dr.rectangle([PAD, y, W - PAD, y + HH], fill="#1F3864")
        x = W - PAD
        for name, key, w, al in COLS:
            cx = x - w / 2 if al == "c" else x - 8
            dr.text((cx, y + HH / 2), get_display(name), font=fh, fill="white", anchor="mm" if al == "c" else "rm")
            x -= w
        y += HH
        for i, r in enumerate(chunk):
            if i % 2 == 1:
                dr.rectangle([PAD, y, W - PAD, y + RH], fill="#F2F5FA")
            x = W - PAD
            for name, key, w, al in COLS:
                v = flag_of(r) if key == "flag" else r.get(key)
                s = _fmt(key, v)
                color = "black"
                if key in ("z", "rel_z") and isinstance(v, (int, float)) and v >= 2:
                    color = "#C00000"
                if key == "dev_bp":
                    color = "#C00000"
                fnt = fb if key == "name" else f
                if al == "c":
                    dr.text((x - w / 2, y + RH / 2), get_display(s), font=fnt, fill=color, anchor="mm")
                else:
                    dr.text((x - 8, y + RH / 2), get_display(s), font=fnt, fill=color, anchor="rm")
                x -= w
            dr.line([PAD, y + RH, W - PAD, y + RH], fill="#D9D9D9")
            y += RH
        p = f"{out_prefix}_{ci+1}.png" if len(chunks) > 1 else f"{out_prefix}.png"
        img.save(p); paths.append(p)
    return paths


def render_xlsx(df, path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    heb = {"name": "סדרה", "issuer": "מנפיק", "covered": "מסוקר", "rating_low": "דירוג", "linkage": "הצמדה", "duration": "מח\"מ", "cur_date": "תאריך",
           "cur_spread": "מרווח %", "mean": "ממוצע 3 ש' %", "std": "סטיית תקן", "min": "מינ' %", "max": "מקס' %", "dev_bp": "סטייה (נ\"ב)", "z": "Z",
           "rel_z": "Z יחסי", "rel_dev_bp": "סטייה יחסית (נ\"ב)", "pct_rank": "אחוזון", "chg_5d_bp": "5 ימים (נ\"ב)", "chg_vs_lastq_bp": "מול רבעון (נ\"ב)",
           "peer_cur_median": "חציון קבוצה %", "rating_change": "שינוי דירוג", "maya_recent": "דיווח מאיה", "note": "הערה", "sent_today": "נשלח היום", "security_id": "מס' ני\"ע"}
    fmt = {"duration": "0.00", "cur_spread": "0.00", "mean": "0.00", "std": "0.00", "min": "0.00", "max": "0.00", "dev_bp": "0", "z": "0.0", "rel_z": "0.0",
           "rel_dev_bp": "0", "pct_rank": "0%", "chg_5d_bp": "0", "chg_vs_lastq_bp": "0", "peer_cur_median": "0.00"}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "מרווחים חריגים"; ws.sheet_view.rightToLeft = True
    cols = list(df.columns)
    for j, c in enumerate(cols):
        cell = ws.cell(1, j + 1, heb.get(c, c)); cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(j + 1)].width = 22 if c in ("name", "issuer", "note") else 10
    for i, (_, r) in enumerate(df.iterrows()):
        for j, c in enumerate(cols):
            v = r[c]; v = None if (v is None or (isinstance(v, float) and v != v)) else v
            if c == "covered": v = "★" if isinstance(v, str) and v else None
            if c == "linkage": v = LINK.get(v, v)
            cell = ws.cell(i + 2, j + 1, v); cell.font = Font(name="Arial", size=10)
            if c in fmt: cell.number_format = fmt[c]
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{L(len(cols))}{max(2, len(df) + 1)}"; ws.row_dimensions[1].height = 30
    wb.save(path)
