"""Three-way, cell-by-cell merge of an Excel model (used when Yovav's PC and the cloud both changed a file).

    python corp/merge_xlsx.py BASE.xlsx OURS.xlsx THEIRS.xlsx OUT.xlsx [--report report.json]

BASE   = the version both sides had at the last sync (git blob `base_blob` in _sync/conflicts/*.json)
OURS   = Yovav's version (the <name>_מקומי_<time>.xlsx copy)
THEIRS = the current cloud version (the bot's update)

OUT starts as THEIRS and gets every cell Yovav changed that the cloud did not touch. A cell changed on both
sides to different values is a conflict: THEIRS is kept and the cell is listed. Sheets Yovav added are
copied (values/formulas + number formats); sheets he deleted are only reported. Formulas compare as text.
"""
import argparse, json, sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def norm(v):
    if v is None or v == "":
        return None
    t = getattr(v, "text", None)          # ArrayFormula / DataTableFormula
    return f"={t}" if t is not None and not isinstance(v, str) else v


def cells(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            v = norm(c.value)
            if v is not None:
                out[c.coordinate] = v
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("base"); ap.add_argument("ours"); ap.add_argument("theirs"); ap.add_argument("out")
    ap.add_argument("--report", default="")
    a = ap.parse_args()
    wb_b = load_workbook(a.base)
    wb_o = load_workbook(a.ours)
    wb_t = load_workbook(a.theirs)
    out = load_workbook(a.theirs)
    rep = {"applied": [], "conflicts": [], "new_sheets": [], "removed_sheets": [], "skipped_sheets": []}

    for name in wb_o.sheetnames:
        wo = wb_o[name]
        if name not in wb_b.sheetnames:
            if name in out.sheetnames:
                rep["skipped_sheets"].append(name)
                continue
            ws = out.create_sheet(name)
            ws.sheet_view.rightToLeft = wo.sheet_view.rightToLeft
            for coord, v in cells(wo).items():
                ws[coord] = v
                ws[coord].number_format = wo[coord].number_format
            for k, d in wo.column_dimensions.items():
                ws.column_dimensions[k].width = d.width
            rep["new_sheets"].append(name)
            continue
        if name not in wb_t.sheetnames:
            rep["skipped_sheets"].append(name)
            continue
        cb, co, ct = cells(wb_b[name]), cells(wo), cells(wb_t[name])
        ws = out[name]
        for coord in sorted(set(cb) | set(co) | set(ct)):
            b, o, t = cb.get(coord), co.get(coord), ct.get(coord)
            if o == b:
                continue
            if t == b:
                ws[coord] = o
                if o is not None and wo[coord].number_format != wb_b[name][coord].number_format:
                    ws[coord].number_format = wo[coord].number_format
                rep["applied"].append({"cell": f"{name}!{coord}", "base": repr(b), "ours": repr(o)})
            elif o != t:
                rep["conflicts"].append({"cell": f"{name}!{coord}", "base": repr(b), "ours": repr(o), "theirs": repr(t)})
    for name in wb_b.sheetnames:
        if name not in wb_o.sheetnames:
            rep["removed_sheets"].append(name)

    out.save(a.out)
    rep["summary"] = (f"applied {len(rep['applied'])} cells, {len(rep['conflicts'])} conflicts, "
                      f"new sheets {rep['new_sheets']}, removed sheets {rep['removed_sheets']}")
    if a.report:
        with open(a.report, "w", encoding="utf-8") as fh:
            json.dump(rep, fh, ensure_ascii=False, indent=1)
    print(rep["summary"])
    for c in rep["conflicts"][:30]:
        print("  CONFLICT", c["cell"], "base", c["base"], "| ours", c["ours"], "| theirs", c["theirs"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
