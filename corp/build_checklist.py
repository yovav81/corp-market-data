#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_checklist.py - רשימת הבדיקות לדוח הבא, לכל המנפיקים.
אוסף מכל מודל את טבלת ׳שאלה / פריט חסר׳ בגיליון מקורות (רק פריטים שלא נסגרו) + פריטים ידניים מ-00_Master/checks_extra.csv.
מ-29.10.2026 הרשימה היא גיליון בתוך אקסל העל (build_live.py קורא ל-add_sheets); הקובץ הנפרד אינו נבנה יותר.
שימוש עצמאי (קובץ נפרד, לבדיקה): python3 00_Master/_build/build_checklist.py 02_Issuers 00_Master/checks_extra.csv <פלט.xlsx>"""
import sys, glob, os, csv, datetime as dt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

SEV = {"גבוהה": 0, "בינונית": 1, "נמוכה": 2}
CLOSED = ("נסגר", "סגור")
SHEET, SUMMARY = "בדיקות_לדוח_הבא", "בדיקות_סיכום"

def pname(wb, n):
    try:
        sh, ref = wb.defined_names[n].attr_text.split("!"); return wb[sh.strip("'")][ref.replace("$", "")].value
    except Exception: return None

def fmt_date(v):
    if isinstance(v, dt.datetime): return v.strftime("%d.%m.%Y")
    return str(v or "")

def collect(iss_dir, extra):
    """כל הפריטים הפתוחים: [מנפיק, מס׳ חברה, סטטוס מודל, תקופה, פריט, גיליון, חומרה, נפתח, פעולה, סטטוס, מקור]"""
    rows, names = [], {}
    for d in sorted(glob.glob(os.path.join(iss_dir, "*_*"))):
        if not os.path.isdir(d): continue
        fs = [f for f in glob.glob(os.path.join(d, "*.xlsx")) if "_מקומי_" not in f and "~$" not in f]
        if not fs: continue
        f = max(fs, key=os.path.getmtime)
        folder = os.path.basename(d); iid, nm = folder.split("_", 1); names[iid] = nm
        wb = openpyxl.load_workbook(f, data_only=True)
        st = "מאושר" if str(pname(wb, "p_Analyst") or "").startswith("מאושר") else "טיוטה"
        per = pname(wb, "p_Period") or ""
        if "מקורות" not in wb.sheetnames: continue
        m = wb["מקורות"]; hdr = None
        for r in range(1, m.max_row + 1):
            if m.cell(r, 2).value == "שאלה / פריט חסר": hdr = r; break
        if not hdr: continue
        for r in range(hdr + 1, m.max_row + 1):
            item = m.cell(r, 2).value
            if not item: continue
            status = str(m.cell(r, 7).value or "")
            act = str(m.cell(r, 6).value or "").strip(); sev = str(m.cell(r, 4).value or "").strip()
            if status.startswith(CLOSED) or status == "טופל" or act.startswith(CLOSED) or sev == "סגור" or act.startswith("ללא פעולה"): continue
            rows.append([nm, iid, st, per, item, m.cell(r, 3).value, m.cell(r, 4).value, fmt_date(m.cell(r, 5).value), m.cell(r, 6).value, status or "פתוח", f"מודל — מקורות!B{r}"])
    if extra and os.path.exists(extra):
        with open(extra, encoding="utf-8-sig") as fh:
            for e in csv.DictReader(fh):
                iid = e["issuer_id"]
                if e["status"].startswith(CLOSED): continue
                rows.append([names.get(iid, "כללי" if iid == "כללי" else iid), iid, "", "", e["item"], e["sheet"], e["severity"], e["opened"], e["action"], e["status"], "checks_extra.csv"])
    rows.sort(key=lambda x: (x[0] == "כללי", x[0], SEV.get(x[6], 3)))
    return rows

def add_sheets(wb, iss_dir, extra, index=None):
    """מוסיף לחוברת wb את הגיליונות בדיקות_לדוח_הבא + בדיקות_סיכום (במיקום index). מחזיר (מס׳ פריטים, מס׳ מנפיקים)."""
    rows = collect(iss_dir, extra)
    for n in (SHEET, SUMMARY):
        if n in wb.sheetnames: del wb[n]
    ws = wb.create_sheet(SHEET, index); ws.sheet_view.rightToLeft = True
    ws["A1"] = "בדיקות לדוח הבא — כל המנפיקים"; ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = (f"נבנה אוטומטית {dt.date.today():%d.%m.%Y} מטבלאות ׳שאלה / פריט חסר׳ בגיליון מקורות של כל מודל (פריטים שלא נסגרו; בלי פריטים שסומנו ׳ללא פעולה׳) ומ-00_Master/checks_extra.csv. "
                "בכל קליטת דוח: לעבור על השורות של המנפיק, לטפל, ולסמן ׳נסגר׳ במודל — הרשימה מתעדכנת בבנייה הבאה של אקסל העל.")
    ws["A2"].font = Font(name="Arial", size=10); ws["A2"].alignment = Alignment(wrap_text=True); ws.merge_cells("A2:K2"); ws.row_dimensions[2].height = 32
    H = ["מנפיק", "מס׳ חברה", "סטטוס מודל", "תקופה", "פריט לבדיקה", "גיליון", "חומרה", "נפתח", "פעולה נדרשת", "סטטוס", "מקור"]
    ws.append([]); ws.append(H)
    for x in rows: ws.append(x)
    W = [16, 11, 10, 9, 70, 18, 9, 11, 50, 18, 18]
    for i, w in enumerate(W): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    fill = {"גבוהה": "F8D7DA", "בינונית": "FFF3CD", "נמוכה": "E2EFDA"}
    for r in range(5, ws.max_row + 1):
        for c in range(1, 12):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(r, c).font = Font(name="Arial", size=10)
        s = ws.cell(r, 7).value
        if s in fill: ws.cell(r, 7).fill = PatternFill("solid", fgColor=fill[s])
    t = Table(displayName="tblChecks", ref=f"A4:K{max(5, ws.max_row)}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A5"
    s2 = wb.create_sheet(SUMMARY, (wb.index(ws) + 1)); s2.sheet_view.rightToLeft = True
    s2["A1"] = "בדיקות לדוח הבא — סיכום לפי מנפיק"; s2["A1"].font = Font(name="Arial", bold=True, size=14)
    s2.append([]); s2.append(["מנפיק", "פתוחים", "גבוהה", "בינונית", "נמוכה"])
    for c in range(1, 6): s2.cell(3, c).font = Font(name="Arial", bold=True)
    agg = {}
    for x in rows:
        a = agg.setdefault(x[0], [0, 0, 0, 0]); a[0] += 1
        if x[6] in SEV: a[1 + SEV[x[6]]] += 1
    for k, v in agg.items(): s2.append([k] + v)
    s2.column_dimensions["A"].width = 20; s2.freeze_panes = "A4"
    return len(rows), len(agg)

if __name__ == "__main__":
    iss_dir, extra, out = sys.argv[1:4]
    wb = openpyxl.Workbook(); del wb[wb.active.title]
    n, k = add_sheets(wb, iss_dir, extra)
    wb.save(out); print("saved", out, n, "items,", k, "issuers")
