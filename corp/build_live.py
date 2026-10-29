#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_live.py — בונה את אקסל העל היחיד: 00_Master_LIVE.xlsx (מ-29.10.2026 אין יותר גרסאות 00_Master_vNN).

מה הקובץ מקבל:
 1. כל הגיליונות של build_master.py (+ גורמים_חיצוניים מ-master_factors.py) — תמונת מצב מהמודלים ומנתוני השוק בענן.
 2. גיליונות בדיקות_לדוח_הבא + בדיקות_סיכום (build_checklist.add_sheets) — במקום הקובץ הנפרד.
 3. קלט המשתמש מהקובץ הקודם נשמר: Watchlist (מס׳ ני״ע + הערה), ערכי פרמטרים (לפי שם הפרמטר),
    דריסות ידניות של מדדים צפויים ב׳מחמ מתחת ל 1׳ (לפי חודש).
 4. שלוש שאילתות Power Query (Issuers / Series / Market) — קוראות את C:\\CORP\\02_Issuers ו-C:\\CORP\\03_Market\\data
    ומרעננות בפתיחת הקובץ, כך שבמחשב הקובץ מתעדכן גם בין בנייה לבנייה. ה-DataMashup נלקח מהקובץ הקודם
    (כדי ששינוי שיובב עושה בשאילתות באקסל יישמר), ואם אין — מ-_build/pq_kit.

שימוש:
  python3 00_Master/_build/build_live.py <corp-data> <corp-market-data> [--out <xlsx>] [--prev <xlsx>] [--check]
  ברירת מחדל: --out <corp-data>/00_Master_LIVE.xlsx ; --prev = הקובץ הקיים באותו נתיב.
  --check: חישוב מחדש ב-LibreOffice על עותק (לפני הזרקת השאילתות) וספירת שגיאות נוסחה.
"""
import sys, os, re, io, glob, shutil, subprocess, tempfile, zipfile, datetime as dt
import xml.etree.ElementTree as ET
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl
from openpyxl.utils import get_column_letter as L

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
# שם גיליון → (שם הטבלה/שאילתה, מזהה החיבור ב-connections.xml)
QUERIES = {"נתוני_מנפיקים": ("Issuers", 1), "נתוני_שוק": ("Market", 2), "נתוני_סדרות": ("Series", 3)}
LIVE_NAME = "00_Master_LIVE.xlsx"


def run(cmd, **kw):
    print("$", " ".join(str(c) for c in cmd)); subprocess.run([str(c) for c in cmd], check=True, **kw)


# ---------------------------------------------------------------- 1+2: בנייה
def build_snapshot(cd, cmd, out):
    sp = os.path.join(cmd, "data", "spreads", f"{dt.date.today().year}.csv")
    if not os.path.exists(sp): sp = sorted(glob.glob(os.path.join(cmd, "data", "spreads", "*.csv")))[-1]
    iss = os.path.join(cd, "02_Issuers"); sec = os.path.join(cd, "00_Master", "sectors.csv"); upd = os.path.join(cd, "00_Master", "updates.csv")
    args = [sys.executable, os.path.join(HERE, "build_master.py"), iss, sp, out, sec]
    if os.path.exists(upd) and _has_header(upd): args.append(upd)
    run(args)
    run([sys.executable, os.path.join(HERE, "master_factors.py"), out, iss, os.path.join(cd, "00_Master", "פרמטרים_רוחביים.xlsx")])
    import build_checklist
    wb = openpyxl.load_workbook(out)
    n, k = build_checklist.add_sheets(wb, iss, os.path.join(cd, "00_Master", "checks_extra.csv"), index=wb.index(wb["מה_חדש"]) + 1)
    print("בדיקות לדוח הבא:", n, "פריטים,", k, "מנפיקים")
    _fix_guide(wb)
    reorder(wb)
    return wb


ORDER = ["השוואה_מנפיקים", "השוואה_ענפית", "Watchlist", "השוואה_אגחים", "סינון_שוק", "מרווחים_חריגים", "מחמ מתחת ל 1", "תנאי_אשראי",
         "מה_חדש", "בדיקות_לדוח_הבא", "בדיקות_סיכום", "גורמים_חיצוניים"]   # סדר הגיליונות שיובב ביקש (29.10); השאר אחריהם בסדר הקיים


def reorder(wb):
    rest = [n for n in wb.sheetnames if n not in ORDER]
    want = [n for n in ORDER if n in wb.sheetnames] + rest
    wb._sheets = [wb[n] for n in want]
    wb.active = 0


def _has_header(p):
    with open(p, encoding="utf-8-sig") as fh: return fh.readline().startswith("date,")


def _fix_guide(wb):
    """מעדכן את שורות הפתיחה במדריך: קובץ חי יחיד, לא גרסאות."""
    if "מדריך" not in wb.sheetnames: return
    g = wb["מדריך"]
    for r in range(1, 6):
        v = g.cell(r, 1).value
        if isinstance(v, str) and v.startswith("אקסל על v1"): g.cell(r, 1).value = "אקסל על — 00_Master_LIVE.xlsx (קובץ יחיד) — מדריך"
        if isinstance(v, str) and v.startswith("גרסה: v1"):
            g.cell(r, 1).value = (f"נבנה בענן {dt.date.today():%d.%m.%Y} (build_live.py) מהמודלים ומנתוני השוק. "
                                  "שלוש השאילתות (Issuers/Series/Market) מרעננות בפתיחה את גיליוני נתוני_* מהקבצים במחשב; "
                                  "Watchlist, פרמטרים ודריסות מדדים נשמרים בין בנייה לבנייה. אין יותר קבצי 00_Master_vNN.")


# ---------------------------------------------------------------- 3: קלט משתמש מהקובץ הקודם
def carry_over(wb, prev_path):
    if not prev_path or not os.path.exists(prev_path): print("אין קובץ קודם — בלי העברת קלט"); return
    import warnings; warnings.simplefilter("ignore")
    pv = openpyxl.load_workbook(prev_path)
    # Watchlist
    if "Watchlist" in pv.sheetnames and "Watchlist" in wb.sheetnames:
        o, n = pv["Watchlist"], wb["Watchlist"]
        rows = [(o.cell(r, 1).value, o.cell(r, 2).value) for r in range(5, 65) if o.cell(r, 1).value not in (None, "")]
        real = [x for x in rows if not (isinstance(x[1], str) and x[1].startswith("דוגמה"))]
        if real:
            for k in range(60): n.cell(5 + k, 1).value = None; n.cell(5 + k, 2).value = None
            for k, (sid, note) in enumerate(rows[:60]): n.cell(5 + k, 1).value = sid; n.cell(5 + k, 2).value = note
            print("Watchlist: הועברו", len(rows), "שורות")
    # פרמטרים — לפי שם הפרמטר
    if "פרמטרים" in pv.sheetnames and "פרמטרים" in wb.sheetnames:
        o, n = pv["פרמטרים"], wb["פרמטרים"]; old = {o.cell(r, 1).value: o.cell(r, 2).value for r in range(5, 40) if o.cell(r, 1).value}
        ch = 0
        for r in range(5, 40):
            lab = n.cell(r, 1).value
            if lab in old and old[lab] is not None and old[lab] != n.cell(r, 2).value: n.cell(r, 2).value = old[lab]; ch += 1
        if ch: print("פרמטרים: הועברו", ch, "ערכים שונים מברירת המחדל")
    # דריסות מדדים (עמודה F) בגיליון מחמ מתחת ל 1 — לפי חודש בעמודה A
    sn = "מחמ מתחת ל 1"
    if sn in pv.sheetnames and sn in wb.sheetnames:
        o, n = pv[sn], wb[sn]
        def idx_rows(ws):
            hdr = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "חודש המדד"), None)
            out = {}
            if hdr:
                r = hdr + 1
                while isinstance(ws.cell(r, 1).value, (dt.datetime, dt.date)): out[ws.cell(r, 1).value.strftime("%Y-%m")] = r; r += 1
            return out
        ov = {k: o.cell(r, 6).value for k, r in idx_rows(o).items() if o.cell(r, 6).value not in (None, "")}
        ch = 0
        for k, r in idx_rows(n).items():
            if k in ov and n.cell(r, 6).fill is not None and n.cell(r, 6).fill.fgColor is not None and str(n.cell(r, 6).fill.fgColor.rgb).endswith("FFF2CC"):
                n.cell(r, 6).value = ov[k]; ch += 1
        if ch: print("מחמ מתחת ל 1: הועברו", ch, "דריסות מדד")


# ---------------------------------------------------------------- 4: הזרקת Power Query (ברמת ה-zip)
def extract_pq(prev_path):
    """מחזיר (DataMashup bytes, itemProps bytes, connections bytes) מהקובץ הקודם, או מ-pq_kit."""
    kit = os.path.join(HERE, "pq_kit")
    if prev_path and os.path.exists(prev_path):
        try:
            z = zipfile.ZipFile(prev_path); names = set(z.namelist())
            cx = [n for n in names if re.fullmatch(r"customXml/item\d+\.xml", n) and _is_mashup(z.read(n)[:600])]
            if cx and "xl/connections.xml" in names:
                props = cx[0].replace("item", "itemProps")
                print("Power Query: נלקח מהקובץ הקודם", os.path.basename(prev_path))
                return z.read(cx[0]), z.read(props) if props in names else open(os.path.join(kit, "itemProps1.xml"), "rb").read(), z.read("xl/connections.xml")
        except Exception as e: print("אזהרה: לא נקראו השאילתות מהקובץ הקודם:", e)
    print("Power Query: נלקח מ-_build/pq_kit")
    return (open(os.path.join(kit, "DataMashup.xml"), "rb").read(), open(os.path.join(kit, "itemProps1.xml"), "rb").read(),
            open(os.path.join(kit, "connections.xml"), "rb").read())


def _is_mashup(b):
    return b"DataMashup" in b or b"D\x00a\x00t\x00a\x00M\x00a\x00s\x00h\x00u\x00p" in b


def inject_pq(xlsx_in, xlsx_out, mashup, props, connections):
    zin = zipfile.ZipFile(xlsx_in); files = {n: zin.read(n) for n in zin.namelist()}
    wb = ET.fromstring(files["xl/workbook.xml"])
    wbrels = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    rid_target = {r.get("Id"): r.get("Target") for r in wbrels}
    sheets = wb.find(f"{{{NS}}}sheets")
    sheet_list = [(s.get("name"), rid_target[s.get(f"{{{RNS}}}id")]) for s in sheets]
    conn_ids = {c.get("name").split("- ", 1)[-1].strip(): c.get("id") for c in ET.fromstring(connections)}  # "שאילתה - Issuers" → "1"
    qt_n = 0; new_defined = []
    for idx, (sname, target) in enumerate(sheet_list):
        if sname not in QUERIES: continue
        qname, default_cid = QUERIES[sname]; cid = conn_ids.get(qname, str(default_cid))
        spath = "xl/" + target.lstrip("/").replace("xl/", "", 1) if not target.startswith("/") else target.lstrip("/")
        srels_path = spath.replace("worksheets/", "worksheets/_rels/") + ".rels"
        srels = ET.fromstring(files[srels_path])
        tpath = None
        for r in srels:
            if r.get("Type").endswith("/table"):
                t = r.get("Target"); tpath = t.lstrip("/") if t.startswith("/") else "xl/" + t.replace("../", "")
        if not tpath: raise SystemExit(f"אין טבלה בגיליון {sname}")
        tb = ET.fromstring(files[tpath])
        tb.set("name", qname); tb.set("displayName", qname); tb.set("tableType", "queryTable"); tb.set("totalsRowShown", "0")
        cols = tb.find(f"{{{NS}}}tableColumns")
        for i, c in enumerate(cols, 1):
            c.set("id", str(i)); c.set("uniqueName", str(i)); c.set("queryTableFieldId", str(i))
        tsi = tb.find(f"{{{NS}}}tableStyleInfo")
        if tsi is None: tsi = ET.SubElement(tb, f"{{{NS}}}tableStyleInfo")
        tsi.set("name", "TableStyleLight1"); tsi.set("showFirstColumn", "0"); tsi.set("showLastColumn", "0"); tsi.set("showRowStripes", "0"); tsi.set("showColumnStripes", "0")
        files[tpath] = xml_bytes(tb)
        # queryTable part
        qt_n += 1; qpath = f"xl/queryTables/queryTable{qt_n}.xml"
        fields = "".join(f'<queryTableField id="{i}" name="{_esc(c.get("name"))}" tableColumnId="{i}"/>' for i, c in enumerate(cols, 1))
        files[qpath] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<queryTable xmlns="{NS}" name="ExternalData_{qt_n}" backgroundRefresh="0" refreshOnLoad="1" '
                        f'connectionId="{cid}" autoFormatId="16" applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" '
                        f'applyAlignmentFormats="0" applyWidthHeightFormats="0"><queryTableRefresh nextId="{len(cols) + 1}"><queryTableFields count="{len(cols)}">'
                        f'{fields}</queryTableFields></queryTableRefresh></queryTable>').encode("utf-8")
        trels = tpath.replace("tables/", "tables/_rels/") + ".rels"
        files[trels] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="{PNS}"><Relationship Id="rId1" '
                        f'Type="{RNS}/queryTable" Target="../queryTables/queryTable{qt_n}.xml"/></Relationships>').encode("utf-8")
        new_defined.append((idx, sname, tb.get("ref"), qt_n))
    # connections + customXml
    files["xl/connections.xml"] = connections
    files["customXml/item1.xml"] = mashup; files["customXml/itemProps1.xml"] = props
    files["customXml/_rels/item1.xml.rels"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="{PNS}"><Relationship Id="rId1" '
                                               f'Type="{RNS}/customXmlProps" Target="itemProps1.xml"/></Relationships>').encode("utf-8")
    nxt = max(int(re.sub(r"\D", "", r.get("Id")) or 0) for r in wbrels) + 1
    for i, (typ, tgt) in enumerate(((f"{RNS}/connections", "connections.xml"), (f"{RNS}/customXml", "../customXml/item1.xml"))):
        e = ET.SubElement(wbrels, f"{{{PNS}}}Relationship"); e.set("Id", f"rId{nxt + i}"); e.set("Type", typ); e.set("Target", tgt)
    files["xl/_rels/workbook.xml.rels"] = xml_bytes(wbrels)
    # defined names ExternalData_n (מוסתרים, ברמת הגיליון)
    dn = wb.find(f"{{{NS}}}definedNames")
    if dn is None:
        dn = ET.Element(f"{{{NS}}}definedNames"); wb.insert(list(wb).index(sheets) + 1, dn)
    for idx, sname, ref, n in new_defined:
        e = ET.SubElement(dn, f"{{{NS}}}definedName"); e.set("name", f"ExternalData_{n}"); e.set("localSheetId", str(idx)); e.set("hidden", "1")
        e.text = f"'{sname}'!{_abs(ref)}"
    files["xl/workbook.xml"] = xml_bytes(wb)
    # content types
    ct = ET.fromstring(files["[Content_Types].xml"])
    have = {o.get("PartName") for o in ct if o.tag.endswith("Override")}
    adds = [("/xl/connections.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml"),
            ("/customXml/itemProps1.xml", "application/vnd.openxmlformats-officedocument.customXmlProperties+xml")]
    adds += [(f"/xl/queryTables/queryTable{i}.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml") for i in range(1, qt_n + 1)]
    for pn, typ in adds:
        if pn not in have:
            e = ET.SubElement(ct, f"{{{CT}}}Override"); e.set("PartName", pn); e.set("ContentType", typ)
    files["[Content_Types].xml"] = xml_bytes(ct)
    with zipfile.ZipFile(xlsx_out, "w", zipfile.ZIP_DEFLATED) as zo:
        zo.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
        for n, b in files.items(): zo.writestr(n, b)
    print("Power Query הוזרק:", qt_n, "שאילתות →", xlsx_out)


def xml_bytes(el):
    """סריאליזציה עם מרחב השמות של האלמנט כברירת מחדל (בלי קידומות ns0)."""
    ns = el.tag[1:].split("}")[0] if el.tag.startswith("{") else ""
    ET.register_namespace("", ns); ET.register_namespace("r", RNS)
    return b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(el, encoding="utf-8")
def _esc(s): return s.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")
def _abs(ref):
    a, b = ref.split(":"); return "$" + re.sub(r"(\d+)", r"$\1", a) + ":$" + re.sub(r"(\d+)", r"$\1", b)


# ---------------------------------------------------------------- בדיקה
def check_errors(xlsx):
    tmp = tempfile.mkdtemp(prefix="live_check_")
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, xlsx], check=True, capture_output=True, timeout=900)
    v = openpyxl.load_workbook(os.path.join(tmp, os.path.basename(xlsx)), data_only=True)
    n = 0; f = 0; per = {}
    for ws in v.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper(): n += 1; per[ws.title] = per.get(ws.title, 0) + 1
    for ws in openpyxl.load_workbook(xlsx).worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="): f += 1
    print("בדיקה: נוסחאות", f, "| שגיאות", n, per or "")
    return n


def main():
    a = sys.argv[1:]
    if len(a) < 2: raise SystemExit(__doc__)
    cd, cmd = os.path.abspath(a[0]), os.path.abspath(a[1])
    out = os.path.abspath(a[a.index("--out") + 1]) if "--out" in a else os.path.join(cd, LIVE_NAME)
    prev = os.path.abspath(a[a.index("--prev") + 1]) if "--prev" in a else (out if os.path.exists(out) else None)
    tmpd = tempfile.mkdtemp(prefix="build_live_"); snap = os.path.join(tmpd, "snapshot.xlsx")
    wb = build_snapshot(cd, cmd, snap)
    carry_over(wb, prev)
    wb.save(snap)
    if "--check" in a and check_errors(snap): print("!! יש שגיאות נוסחה — הקובץ נבנה בכל זאת; לבדוק")
    mashup, props, conns = extract_pq(prev)
    inject_pq(snap, out + ".tmp", mashup, props, conns)
    os.replace(out + ".tmp", out)
    print("saved", out)


if __name__ == "__main__":
    main()
