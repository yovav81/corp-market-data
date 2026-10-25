#!/usr/bin/env python3
"""factors.py — גורמים חיצוניים רוחביים (T1.13): משיכת סדרות אוטומטיות לקובץ הפרמטרים הרוחביים.
python3 factors.py fetch [params.xlsx]   — מושך את כל הסדרות האוטומטיות ומעדכן את גיליונות ׳גורמים׳ ו־׳היסטוריית גורמים׳
python3 factors.py probe out.json        — בדיקת זמינות המקורות (ללא כתיבה)
סדרה ידנית: מוסיפים שורה בגיליון ׳היסטוריית גורמים׳ (קוד, תאריך, ערך, מקור) — הערך האחרון הוא הנוכחי.
"""
import sys, os, io, csv, json, datetime as dt, urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
PARAMS = os.environ.get("CORP_PARAMS", os.path.join(HERE, "..", "פרמטרים_רוחביים.xlsx"))
START = dt.date(2025, 1, 1)
UA = {"User-Agent": "Mozilla/5.0 (corp-data factors; yovav81)"}

# קוד: (שם, קטגוריה, סוג שינוי, יחידה, מקור, פרמטר)
#   סוג שינוי: ׳יחסי׳ = ערך חדש/ערך בסיס − 1 ; ׳נק׳ אחוז׳ = ערך חדש − ערך בסיס (שבר: 0.01 = 1 נק׳)
CATALOG = {
    "BOI":       ("ריבית בנק ישראל", "ריבית", "נק׳ אחוז", "שבר", "params", "ריבית בנק ישראל"),
    "PRIME":     ("ריבית פריים", "ריבית", "נק׳ אחוז", "שבר", "params", "BOI+prime_spread"),
    "SOFR":      ("SOFR (ארה״ב, יומי)", "ריבית", "נק׳ אחוז", "שבר", "fred", "SOFR"),
    "EURIBOR3M": ("יוריבור 3 חודשים", "ריבית", "נק׳ אחוז", "שבר", "ecb", ("FM", "M.U2.EUR.RT.MM.EURIBOR3MD_.HSTA")),
    "WIBOR3M":   ("WIBOR 3 חודשים (פולין)", "ריבית", "נק׳ אחוז", "שבר", "manual", None),
    "SONIA":     ("SONIA (בריטניה)", "ריבית", "נק׳ אחוז", "שבר", "fred", "IUDSOIA"),
    "UST10":     ("תשואת אג״ח ממשלת ארה״ב 10 שנים", "ריבית", "נק׳ אחוז", "שבר", "fred", "DGS10"),
    "CPI_IL":    ("מדד המחירים לצרכן (הלמ״ס)", "מדד", "יחסי", "נקודות מדד", "cbs", 120010),
    "CONST_IL":  ("מדד תשומות הבנייה למגורים (הלמ״ס)", "תשומות בנייה", "יחסי", "נקודות מדד", "cbs", 200010),
    "HOUSE_IL":  ("מדד מחירי הדירות (הלמ״ס) — קירוב למחירי המכירה", "שוק הדיור", "יחסי", "נקודות מדד", "cbs", 40010),
    "USDILS":    ("דולר / שקל (שער יציג)", "מט״ח", "יחסי", "ש״ח", "boi", "RER_USD_ILS"),
    "EURILS":    ("אירו / שקל (שער יציג)", "מט״ח", "יחסי", "ש״ח", "boi", "RER_EUR_ILS"),
    "GBPILS":    ("ליש״ט / שקל (שער יציג)", "מט״ח", "יחסי", "ש״ח", "boi", "RER_GBP_ILS"),
    "CHFILS":    ("פרנק שוויצרי / שקל (שער יציג)", "מט״ח", "יחסי", "ש״ח", "boi", "RER_CHF_ILS"),
    "PLNILS":    ("זלוטי / שקל (שער צולב ECB)", "מט״ח", "יחסי", "ש״ח", "ecbx", "PLN"),
    "BRLILS":    ("ריאל ברזילאי / שקל (שער צולב ECB)", "מט״ח", "יחסי", "ש״ח", "ecbx", "BRL"),
    "BRENT":     ("נפט ברנט", "סחורות", "יחסי", "$ לחבית", "fred", "DCOILBRENTEU"),
    "WTI":       ("נפט WTI", "סחורות", "יחסי", "$ לחבית", "fred", "DCOILWTICO"),
    "TTF":       ("גז טבעי אירופה (TTF, חודשי)", "סחורות", "יחסי", "$ ל-MMBtu", "fred", "PNGASEUUSDM"),
    "HH":        ("גז טבעי ארה״ב (Henry Hub)", "סחורות", "יחסי", "$ ל-MMBtu", "fred", "DHHNGSP"),
    "PL_POWER":  ("מחיר חשמל בפולין (TGE, בסיס)", "מחירי חשמל", "יחסי", "PLN ל-MWh", "manual", None),
    "US_POWER":  ("מחיר חשמל סיטונאי בארה״ב (אזורי הפעילות)", "מחירי חשמל", "יחסי", "$ ל-MWh", "manual", None),
    "IL_GEN":    ("רכיב הייצור בתעריף החשמל (רשות החשמל)", "מחירי חשמל", "יחסי", "אג׳ לקוט״ש", "manual", None),
}


def get(url, timeout=60):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read().decode("utf-8-sig", errors="replace")


def _d(s):
    s = s.strip()
    for f in ("%Y-%m-%d", "%Y-%m", "%d/%m/%Y"):
        try:
            d = dt.datetime.strptime(s[:10] if f == "%Y-%m-%d" else s, f).date()
            if f == "%Y-%m":
                d = (dt.date(d.year + d.month // 12, d.month % 12 + 1, 1) - dt.timedelta(days=1))
            return d
        except ValueError:
            pass
    return None


def _csv_rows(txt, dcol, vcol):
    out = []
    for r in csv.DictReader(io.StringIO(txt)):
        d, v = _d(r.get(dcol, "") or ""), (r.get(vcol) or "").strip()
        try:
            out.append((d, float(v)))
        except ValueError:
            continue
    return [x for x in out if x[0] and x[0] >= START]


def f_fred(sid):
    txt = get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={START}")
    hdr = txt.splitlines()[0].split(",")
    rows = _csv_rows(txt, hdr[0], sid)
    if sid in ("SOFR", "IUDSOIA", "DGS10"):
        rows = [(d, v / 100) for d, v in rows]
    return rows


def f_boi(series):
    txt = get(f"https://edge.boi.org.il/FusionEdgeServer/sdmx/v2/data/dataflow/BOI.STATISTICS/EXR/1.0/{series}?startperiod={START}&format=csv")
    return _csv_rows(txt, "TIME_PERIOD", "OBS_VALUE")


def f_ecb(arg):
    flow, key = arg
    txt = get(f"https://data-api.ecb.europa.eu/service/data/{flow}/{key}?startPeriod={START:%Y-%m}&format=csvdata")
    return [(d, v / 100) for d, v in _csv_rows(txt, "TIME_PERIOD", "OBS_VALUE")]


_ECB_X = {}
def f_ecbx(cur):
    if not _ECB_X:
        txt = get(f"https://data-api.ecb.europa.eu/service/data/EXR/D.ILS+PLN+BRL.EUR.SP00.A?startPeriod={START}&format=csvdata")
        for r in csv.DictReader(io.StringIO(txt)):
            try:
                _ECB_X.setdefault(r["CURRENCY"], {})[_d(r["TIME_PERIOD"])] = float(r["OBS_VALUE"])
            except (ValueError, KeyError):
                pass
    ils, x = _ECB_X.get("ILS", {}), _ECB_X.get(cur, {})
    return sorted((d, round(ils[d] / x[d], 6)) for d in ils if d in x and d >= START)


def f_stooq(sym):
    txt = get(f"https://stooq.com/q/d/l/?s={sym}&i=d")
    return [(d, v / 100) for d, v in _csv_rows(txt, "Date", "Close")]


CBS_LAG = {120010: 1, 200010: 1, 40010: 2}     # חודשים עד הפרסום (ה-15 בחודש) — ערך ׳ידוע׳ במועד


def f_cbs(cid):
    txt = get(f"https://api.cbs.gov.il/index/data/price?id={cid}&format=json&download=false&startPeriod=01-{START.year}&endPeriod=12-{dt.date.today().year}&lang=he")
    js = json.loads(txt); out = []
    def walk(o):
        if isinstance(o, dict):
            if "year" in o and "month" in o and isinstance(o.get("currBase"), dict):
                y, m = int(o["year"]), int(o["month"])
                v = o["currBase"].get("value")
                if v is not None:
                    lag = CBS_LAG.get(cid, 1); mm = m + lag
                    out.append((dt.date(y + (mm - 1) // 12, (mm - 1) % 12 + 1, 15), float(v)))
            for v in o.values(): walk(v)
        elif isinstance(o, list):
            for v in o: walk(v)
    walk(js)
    return sorted(set(out))


FETCH = {"fred": f_fred, "boi": f_boi, "ecb": f_ecb, "ecbx": f_ecbx, "stooq": f_stooq, "cbs": f_cbs}
SRC_NAME = {"fred": "FRED (St. Louis Fed)", "boi": "בנק ישראל — שער יציג (SDMX)", "ecb": "ECB Data Portal", "ecbx": "ECB — שער צולב מול השקל",
            "stooq": "stooq.com", "cbs": "הלמ״ס — API מדדים", "params": "קובץ הפרמטרים — גיליון ריבית", "manual": "ידני"}


def fetch_all():
    res, err = {}, {}
    for code, (_, _, _, _, src, arg) in CATALOG.items():
        if src not in FETCH: continue
        try:
            res[code] = FETCH[src](arg)
            if not res[code]: err[code] = "no rows"
        except Exception as e:
            err[code] = f"{type(e).__name__}: {e}"[:200]
    return res, err


def boi_series(pw):
    hist = sorted((r[0].date() if hasattr(r[0], "date") else r[0], float(r[1])) for r in pw["ריבית בנק ישראל"].iter_rows(min_row=2, values_only=True) if r[0] and r[1] is not None)
    spread = float({r[0]: r[2] for r in pw["פרמטרים"].iter_rows(min_row=2, values_only=True) if r[0]}.get("prime_spread", 0.015))
    return hist, [(d, round(v + spread, 6)) for d, v in hist]


def load_history(pw):
    h = {}
    if "היסטוריית גורמים" in pw.sheetnames:
        for r in pw["היסטוריית גורמים"].iter_rows(min_row=2, values_only=True):
            if r[0] and r[1] is not None and r[2] is not None:
                d = r[1].date() if hasattr(r[1], "date") else r[1]
                h.setdefault(r[0], {})[d] = (float(r[2]), r[3] or "")
    boi, prime = boi_series(pw)
    h["BOI"] = {d: (v, SRC_NAME["params"]) for d, v in boi}
    h["PRIME"] = {d: (v, SRC_NAME["params"]) for d, v in prime}
    return h


def value_at(hist, code, d):
    """הערך האחרון הידוע עד תאריך d (כולל). מדדי הלמ״ס — לפי מועד הפרסום (׳המדד הידוע׳), לא לפי חודש המדד."""
    s = hist.get(code) or {}
    ks = [k for k in s if k <= d]
    if not ks: return None, None
    k = max(ks); return s[k][0], k


def current(hist, code):
    s = hist.get(code) or {}
    if not s: return None, None
    k = max(s); return s[k][0], k


def write(path, fetched, errors):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    pw = openpyxl.load_workbook(path)
    hist = load_history(pw)
    for code, rows in fetched.items():
        src = SRC_NAME[CATALOG[code][4]]
        for d, v in rows:
            hist.setdefault(code, {})[d] = (v, src)
    hdr_f = Font(name="Arial", b=True, color="FFFFFF"); hdr_p = PatternFill("solid", fgColor="1F3864")
    if "היסטוריית גורמים" in pw.sheetnames: del pw["היסטוריית גורמים"]
    wh = pw.create_sheet("היסטוריית גורמים")
    wh.append(["קוד", "תאריך", "ערך", "מקור"])
    for code in CATALOG:
        if CATALOG[code][4] == "params": continue
        for d in sorted(hist.get(code, {})):
            v, s = hist[code][d]; wh.append([code, d, v, s])
            wh.cell(wh.max_row, 2).number_format = "dd/mm/yyyy"
    for c in wh[1]: c.font, c.fill = hdr_f, hdr_p
    wh.column_dimensions["A"].width = 12; wh.column_dimensions["B"].width = 12; wh.column_dimensions["C"].width = 12; wh.column_dimensions["D"].width = 34
    wh.sheet_view.rightToLeft = True; wh.freeze_panes = "A2"
    if "גורמים" in pw.sheetnames: del pw["גורמים"]
    wg = pw.create_sheet("גורמים", 0)
    wg.append(["קוד", "גורם", "קטגוריה", "סוג שינוי", "יחידה", "ערך נוכחי", "תאריך", "31.12.2025", "30.6.2026", "שינוי מ-31.12.25", "שינוי מ-30.6.26", "מקור", "עדכון", "הערה"])
    for code, (name, cat, kind, unit, src, _) in CATALOG.items():
        cur, cd = current(hist, code)
        b1, _ = value_at(hist, code, dt.date(2025, 12, 31)); b2, _ = value_at(hist, code, dt.date(2026, 6, 30))
        def chg(b):
            if cur is None or b in (None, 0): return None
            return round(cur / b - 1, 6) if kind == "יחסי" else round(cur - b, 6)
        note = errors.get(code, "")
        if src == "manual" and cur is None:
            note = "אין ערך — להזין בגיליון ׳היסטוריית גורמים׳ (קוד, תאריך, ערך, מקור) כולל ערך במועד הבסיס"
        wg.append([code, name, cat, kind, unit, cur, cd, b1, b2, chg(b1), chg(b2), SRC_NAME[src],
                   "ידני" if src in ("manual", "params") else "אוטומטי (יומי)", note])
        r = wg.max_row
        wg.cell(r, 7).number_format = "dd/mm/yyyy"
        fmt = "0.00%" if unit == "שבר" else "#,##0.0000"
        for c in (6, 8, 9): wg.cell(r, c).number_format = fmt
        for c in (10, 11): wg.cell(r, c).number_format = "0.00%"
        if src == "manual":
            for c in range(1, 15): wg.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
    for c in wg[1]: c.font, c.fill = hdr_f, hdr_p; c.alignment = Alignment(wrap_text=True, horizontal="center")
    for col, w in zip("ABCDEFGHIJKLMN", [11, 40, 13, 10, 12, 11, 11, 11, 11, 11, 11, 28, 15, 50]):
        wg.column_dimensions[col].width = w
    wg.sheet_view.rightToLeft = True; wg.freeze_panes = "C2"
    wg["P1"] = ("גורמים חיצוניים רוחביים: כל שורה היא גורם שמודל מנפיק אחד לפחות מושפע ממנו לפי הרגישויות שפרסם (גיליון השפעות_חיצוניות). "
                "סדרה אוטומטית: נמשכת כל יום ע״י פעולת GitHub ב-corp-market-data (data/factors) ומסונכרנת לכאן בהחלה. סדרה ידנית (צהוב): מוסיפים שורה ב׳היסטוריית גורמים׳. "
                "ההחלה על כל המודלים והמאסטר: משימה מתוזמנת בכל יום ראשון, או בבקשה (׳תחיל את הפרמטרים הרוחביים׳).")
    wg["P1"].alignment = Alignment(wrap_text=True, vertical="top"); wg.column_dimensions["P"].width = 60
    pw.save(path)
    return {c: current(hist, c) for c in CATALOG}


def to_csv(out_dir, fetched, errors):
    """corp-market-data: data/factors/history.csv (מצטבר — שורות קיימות נשמרות) + latest.csv + errors.json"""
    os.makedirs(out_dir, exist_ok=True)
    hp = os.path.join(out_dir, "history.csv"); h = {}
    if os.path.exists(hp):
        for r in csv.DictReader(open(hp, encoding="utf-8")):
            h[(r["code"], r["date"])] = (r["value"], r["source"])
    for code, rows in fetched.items():
        for d, v in rows:
            h[(code, str(d))] = (repr(round(v, 8)), SRC_NAME[CATALOG[code][4]])
    with open(hp, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(["code", "date", "value", "source"])
        for (c, d) in sorted(h): w.writerow([c, d, *h[(c, d)]])
    with open(os.path.join(out_dir, "latest.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(["code", "name", "date", "value", "source", "error"])
        for code in CATALOG:
            ks = sorted(d for (c, d) in h if c == code)
            w.writerow([code, CATALOG[code][0], ks[-1] if ks else "", h[(code, ks[-1])][0] if ks else "", SRC_NAME[CATALOG[code][4]], errors.get(code, "")])
    json.dump({"run": dt.datetime.utcnow().isoformat(timespec="minutes"), "errors": errors}, open(os.path.join(out_dir, "errors.json"), "w"), ensure_ascii=False, indent=1)


def read_csv(src):
    """history.csv מקובץ מקומי או מ-URL → {code: [(date, value)]}"""
    txt = get(src) if src.startswith("http") else open(src, encoding="utf-8").read()
    out = {}
    for r in csv.DictReader(io.StringIO(txt)):
        if r["code"] in CATALOG and CATALOG[r["code"]][4] not in ("manual", "params"):
            out.setdefault(r["code"], []).append((dt.date.fromisoformat(r["date"]), float(r["value"])))
    return out


HIST_URL = "https://raw.githubusercontent.com/yovav81/corp-market-data/main/data/factors/history.csv"


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "probe":
        res, err = fetch_all()
        out = {c: {"n": len(v), "first": str(v[0]) if v else None, "last": str(v[-1]) if v else None} for c, v in res.items()}
        out["_errors"] = err
        json.dump(out, open(sys.argv[2], "w"), ensure_ascii=False, indent=1, default=str)
        print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    elif cmd == "csv":
        res, err = fetch_all()
        to_csv(sys.argv[2], res, err)
        print("errors:", err); print({c: (len(v), v[-1] if v else None) for c, v in res.items()})
    elif cmd == "sync":
        path = sys.argv[2] if len(sys.argv) > 2 else PARAMS
        res, err = read_csv(sys.argv[3] if len(sys.argv) > 3 else HIST_URL), {}
        cur = write(path, res, err)
        for c, v in cur.items(): print(c, v)
    elif cmd == "fetch":
        path = sys.argv[2] if len(sys.argv) > 2 else PARAMS
        res, err = fetch_all()
        cur = write(path, res, err)
        print("errors:", err)
        for c, v in cur.items(): print(c, v)
